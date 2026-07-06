from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from sqlite3 import Connection

CHUNK_MAX_CHARS = 1800
CHUNK_OVERLAP_CHARS = 200

URL_RE = re.compile(r"https?://[^\s)>\]]+")
DATE_RE = re.compile(r"\b(?:19|20)\d{2}[-/.年](?:0?[1-9]|1[0-2])(?:[-/.月](?:0?[1-9]|[12]\d|3[01])日?)?\b")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)
FORM_FIELD_RE = re.compile(r"^\s*([A-Za-z0-9_\- /一-龥ぁ-んァ-ンー]+?)\s*[:：]\s*(.+?)\s*$", re.MULTILINE)

TECH_KEYWORDS = {
    "ai",
    "api",
    "chroma",
    "docker",
    "excel",
    "fastapi",
    "faiss",
    "fts5",
    "github",
    "lancedb",
    "llama",
    "markdown",
    "next.js",
    "obsidian",
    "ollama",
    "openpyxl",
    "pdf",
    "pypdf",
    "python",
    "qwen",
    "react",
    "sqlite",
    "sqlite fts5",
    "tauri",
    "typescript",
    "vector",
}


@dataclass(frozen=True)
class ExtractedTrace:
    trace_type: str
    value: str
    raw_text: str
    confidence: float
    extractor: str = "rule"


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_openpyxl_text(path)
    if suffix == ".xls":
        return read_xls_text(path)
    if suffix == ".pdf":
        return read_pdf_text(path)
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp932", "shift_jis"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_openpyxl_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel extraction requires openpyxl. Install project dependencies.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for sheet in workbook.worksheets:
            lines = [f"# Sheet: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                values = [format_cell(value) for value in row]
                values = [value for value in values if value]
                if values:
                    lines.append("\t".join(values))
            if len(lines) > 1:
                sections.append("\n".join(lines))
        return "\n\n".join(sections)
    finally:
        workbook.close()


def read_xls_text(path: Path) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("XLS extraction requires xlrd. Install project dependencies.") from exc

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sections: list[str] = []
        for sheet in workbook.sheets():
            lines = [f"# Sheet: {sheet.name}"]
            for row_index in range(sheet.nrows):
                values = [format_cell(sheet.cell_value(row_index, col_index)) for col_index in range(sheet.ncols)]
                values = [value for value in values if value]
                if values:
                    lines.append("\t".join(values))
            if len(lines) > 1:
                sections.append("\n".join(lines))
        return "\n\n".join(sections)
    finally:
        workbook.release_resources()


def format_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires pypdf. Install project dependencies.") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            pages.append(f"# Page {index}\n{text}")
    return "\n\n".join(pages)


def hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def estimate_tokens(text: str) -> int:
    # A rough local estimate is enough for Phase 2 chunk bookkeeping.
    return max(1, len(text) // 4)


def split_chunks(text: str) -> list[tuple[str, str | None]]:
    paragraphs = re.split(r"\n\s*\n", text.strip())
    chunks: list[tuple[str, str | None]] = []
    current: list[str] = []
    current_hint: str | None = None

    def flush() -> None:
        nonlocal current, current_hint
        content = "\n\n".join(part for part in current if part).strip()
        if content:
            chunks.append((content, current_hint))
        current = []
        current_hint = None

    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        heading = HEADING_RE.match(paragraph)
        if heading and current:
            flush()
        if heading:
            current_hint = heading.group(2).strip()
        candidate = "\n\n".join([*current, paragraph]).strip()
        if len(candidate) > CHUNK_MAX_CHARS and current:
            flush()
            if len(paragraph) > CHUNK_MAX_CHARS:
                start = 0
                while start < len(paragraph):
                    end = start + CHUNK_MAX_CHARS
                    chunks.append((paragraph[start:end].strip(), current_hint))
                    start = max(end - CHUNK_OVERLAP_CHARS, end)
                current = []
            else:
                current = [paragraph]
        else:
            current.append(paragraph)

    flush()
    return chunks


def extract_traces(text: str, title: str, path: str) -> list[ExtractedTrace]:
    traces: list[ExtractedTrace] = [
        ExtractedTrace("title", title, title, 1.0),
        ExtractedTrace("path", path, path, 1.0),
    ]

    lower_text = text.lower()
    for keyword in sorted(TECH_KEYWORDS, key=len, reverse=True):
        if keyword in lower_text:
            traces.append(ExtractedTrace("technology", keyword, keyword, 0.8))

    for match in HEADING_RE.finditer(text):
        value = match.group(2).strip()
        if value:
            traces.append(ExtractedTrace("heading", value, match.group(0), 0.9))

    for match in URL_RE.finditer(text):
        traces.append(ExtractedTrace("url", match.group(0), match.group(0), 0.95))

    for match in DATE_RE.finditer(text):
        traces.append(ExtractedTrace("date", match.group(0), match.group(0), 0.75))

    for match in FORM_FIELD_RE.finditer(text):
        key = " ".join(match.group(1).split())
        value = match.group(2).strip()
        if 2 <= len(key) <= 60 and value:
            traces.append(ExtractedTrace("field", key, match.group(0), 0.65))

    deduped: dict[tuple[str, str], ExtractedTrace] = {}
    for trace in traces:
        key = (trace.trace_type, trace.value.lower())
        deduped.setdefault(key, trace)
    return list(deduped.values())


def reindex_document(conn: Connection, document_id: int, absolute_path: str, title: str, relative_path: str) -> None:
    path = Path(absolute_path)
    text = read_text(path)
    chunks = split_chunks(text)
    traces = extract_traces(text, title, relative_path)

    conn.execute("DELETE FROM chunks_fts WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM traces_fts WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM traces WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))

    for chunk_index, (content, location_hint) in enumerate(chunks):
        cursor = conn.execute(
            """
            INSERT INTO chunks (document_id, chunk_index, content, token_count, location_hint, content_hash)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (document_id, chunk_index, content, estimate_tokens(content), location_hint, hash_text(content)),
        )
        chunk_id = int(cursor.lastrowid)
        conn.execute(
            """
            INSERT INTO chunks_fts (content, title, path, document_id, chunk_id)
            VALUES (?, ?, ?, ?, ?)
            """,
            (content, title, relative_path, document_id, chunk_id),
        )

    for trace in traces:
        cursor = conn.execute(
            """
            INSERT INTO traces (document_id, chunk_id, trace_type, value, raw_text, confidence, extractor)
            VALUES (?, NULL, ?, ?, ?, ?, ?)
            """,
            (
                document_id,
                trace.trace_type,
                trace.value,
                trace.raw_text,
                trace.confidence,
                trace.extractor,
            ),
        )
        trace_id = int(cursor.lastrowid)
        conn.execute(
            """
            INSERT INTO traces_fts (value, raw_text, trace_type, path, document_id, trace_id)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (trace.value, trace.raw_text, trace.trace_type, relative_path, document_id, trace_id),
        )

    conn.execute(
        """
        UPDATE documents
        SET extraction_status = 'success', error_message = NULL
        WHERE id = ?
        """,
        (document_id,),
    )


def clear_document_index(conn: Connection, document_id: int) -> None:
    conn.execute("DELETE FROM chunks_fts WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM traces_fts WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM traces WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
